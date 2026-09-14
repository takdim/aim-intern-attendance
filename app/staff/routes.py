from datetime import date, datetime, timedelta
import io

from flask import Blueprint, Response, flash, redirect, render_template, request, url_for
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.extensions import db
from app.models import Attendance, AttendanceSetting, LeaveRequest, Role, User
from app.utils.audit import log_action
from app.utils.decorators import role_required
from app.utils.time_utils import now_in_tz


staff_bp = Blueprint("staff", __name__, url_prefix="/staff")


def _get_today_in_system_tz() -> date:
    setting = AttendanceSetting.get_active()
    tz_name = setting.timezone if setting else "Asia/Makassar"
    return now_in_tz(tz_name).date()


def _resolve_period(period: str, anchor_date: date) -> tuple[date, date]:
    if period == "week":
        start_date = anchor_date - timedelta(days=anchor_date.weekday())
        end_date = start_date + timedelta(days=6)
        return start_date, end_date

    if period == "month":
        start_date = anchor_date.replace(day=1)
        next_month = (start_date.replace(day=28) + timedelta(days=4)).replace(day=1)
        end_date = next_month - timedelta(days=1)
        return start_date, end_date

    return anchor_date, anchor_date


def _resolve_filter_inputs():
    allowed = {"day", "week", "month"}
    period = request.args.get("period", "day").lower()
    if period not in allowed:
        period = "day"

    raw_anchor = request.args.get("anchor", "").strip()
    if raw_anchor:
        try:
            anchor_date = datetime.strptime(raw_anchor, "%Y-%m-%d").date()
        except ValueError:
            anchor_date = _get_today_in_system_tz()
    else:
        anchor_date = _get_today_in_system_tz()

    start_date, end_date = _resolve_period(period, anchor_date)
    return period, anchor_date, start_date, end_date


@staff_bp.route("/dashboard")
@role_required("STAFF", "ADMIN")
def dashboard():
    today = _get_today_in_system_tz()

    total_intern = User.query.join(Role).filter(Role.name == "INTERN", User.deleted_at.is_(None)).count()
    already_check_in = Attendance.query.filter(Attendance.date == today, Attendance.check_in.isnot(None)).count()
    not_check_in = max(total_intern - already_check_in, 0)
    late = Attendance.query.filter_by(date=today, status="TERLAMBAT").count()
    izin = Attendance.query.filter_by(date=today, status="IZIN").count()
    sakit = Attendance.query.filter_by(date=today, status="SAKIT").count()

    attendance_today = Attendance.query.filter_by(date=today).order_by(Attendance.id.desc()).all()

    return render_template(
        "staff/dashboard.html",
        total_intern=total_intern,
        already_check_in=already_check_in,
        not_check_in=not_check_in,
        late=late,
        izin=izin,
        sakit=sakit,
        attendance_today=attendance_today,
    )


@staff_bp.route("/attendance")
@role_required("STAFF", "ADMIN")
def attendance():
    period, anchor_date, start_date, end_date = _resolve_filter_inputs()

    query = Attendance.query.filter(Attendance.date >= start_date, Attendance.date <= end_date)
    rows = query.order_by(Attendance.date.desc(), Attendance.id.desc()).all()

    return render_template(
        "staff/attendance.html",
        attendances=rows,
        period=period,
        anchor_date=anchor_date,
        start_date=start_date,
        end_date=end_date,
        total_records=len(rows),
        late_count=sum(1 for row in rows if row.status == "TERLAMBAT"),
        alpha_count=sum(1 for row in rows if row.status == "ALPHA"),
        permit_count=sum(1 for row in rows if row.status in {"IZIN", "SAKIT", "PULANG_CEPAT"}),
    )


@staff_bp.route("/leave-requests")
@role_required("STAFF", "ADMIN")
def leave_requests():
    status = request.args.get("status")
    query = LeaveRequest.query

    if status:
        query = query.filter(LeaveRequest.status == status)

    rows = query.order_by(LeaveRequest.id.desc()).all()
    return render_template("staff/leave_requests.html", leave_requests=rows)


@staff_bp.route("/leave-requests/<int:request_id>/review", methods=["POST"])
@role_required("STAFF", "ADMIN")
def review_leave_request(request_id: int):
    row = LeaveRequest.query.get_or_404(request_id)

    decision = request.form.get("decision")
    notes = request.form.get("review_notes", "").strip() or None

    if row.status != "PENDING":
        flash("Pengajuan sudah diproses", "warning")
        return redirect(url_for("staff.leave_requests"))

    row.status = "APPROVED" if decision == "APPROVE" else "REJECTED"
    row.review_notes = notes
    row.reviewed_by = current_user.id
    row.reviewed_at = datetime.utcnow()

    if row.status == "APPROVED" and row.request_type in {"SAKIT", "IZIN"}:
        current_day = row.start_date
        while current_day <= row.end_date:
            attendance = Attendance.query.filter_by(user_id=row.user_id, date=current_day).first()
            if not attendance:
                attendance = Attendance(user_id=row.user_id, date=current_day)
            if not attendance.check_in and not attendance.check_out:
                attendance.status = row.request_type
            db.session.add(attendance)
            current_day = current_day.fromordinal(current_day.toordinal() + 1)

    log_action(
        actor_user_id=current_user.id,
        action="REVIEW_LEAVE_REQUEST",
        entity="leave_requests",
        entity_id=row.id,
        metadata={
            "decision": row.status,
            "request_type": row.request_type,
            "user_id": row.user_id,
        },
    )

    db.session.commit()
    flash("Pengajuan izin diproses", "success")
    return redirect(url_for("staff.leave_requests"))


@staff_bp.route("/reports")
@role_required("STAFF", "ADMIN")
def reports():
    period, anchor_date, start_date, end_date = _resolve_filter_inputs()
    rows = (
        Attendance.query.filter(Attendance.date >= start_date, Attendance.date <= end_date)
        .order_by(Attendance.date.desc(), Attendance.id.desc())
        .all()
    )
    return render_template(
        "staff/reports.html",
        attendances=rows,
        period=period,
        anchor_date=anchor_date,
        start_date=start_date,
        end_date=end_date,
        total_records=len(rows),
        hadir_count=sum(1 for row in rows if row.status == "HADIR"),
        terlambat_count=sum(1 for row in rows if row.status == "TERLAMBAT"),
    )


@staff_bp.route("/reports/export-excel")
@role_required("STAFF", "ADMIN")
def export_reports_excel():
    period, anchor_date, start_date, end_date = _resolve_filter_inputs()
    rows = (
        Attendance.query.filter(Attendance.date >= start_date, Attendance.date <= end_date)
        .order_by(Attendance.date.asc(), Attendance.id.asc())
        .all()
    )

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Absensi"

    ws.merge_cells("A1:G1")
    ws["A1"] = "Laporan Kehadiran Staff"
    ws["A1"].font = Font(name="Poppins", size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = f"Periode: {start_date.isoformat()} s/d {end_date.isoformat()} ({period.upper()})"
    ws["A2"].font = Font(name="Inter", size=11)
    ws["A2"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A3:G3")
    ws["A3"] = f"Tanggal Acuan: {anchor_date.isoformat()}"
    ws["A3"].font = Font(name="Inter", size=10, italic=True)
    ws["A3"].alignment = Alignment(horizontal="center")

    headers = ["No", "Tanggal", "Nama", "Check-in", "Check-out", "Status", "Zona"]
    header_row = 5
    for idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=idx, value=header)
        cell.font = Font(name="Inter", bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for i, row in enumerate(rows, start=1):
        ws.cell(row=header_row + i, column=1, value=i)
        ws.cell(row=header_row + i, column=2, value=row.date.isoformat())
        ws.cell(row=header_row + i, column=3, value=row.user.name if row.user else "-")
        ws.cell(
            row=header_row + i,
            column=4,
            value=row.check_in.strftime("%H:%M") if row.check_in else "-",
        )
        ws.cell(
            row=header_row + i,
            column=5,
            value=row.check_out.strftime("%H:%M") if row.check_out else "-",
        )
        ws.cell(row=header_row + i, column=6, value=row.status)
        ws.cell(row=header_row + i, column=7, value=row.zone.name if row.zone else "-")

    for col, width in {"A": 6, "B": 14, "C": 30, "D": 12, "E": 12, "F": 16, "G": 24}.items():
        ws.column_dimensions[col].width = width

    total_row = header_row + len(rows) + 2
    ws.merge_cells(f"A{total_row}:G{total_row}")
    ws[f"A{total_row}"] = f"Total Record: {len(rows)}"
    ws[f"A{total_row}"].font = Font(name="Inter", bold=True)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    filename = f"laporan_staff_{period}_{start_date.isoformat()}_{end_date.isoformat()}.xlsx"
    return Response(
        output.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
