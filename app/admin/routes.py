import csv
import io
from datetime import date, datetime, timedelta

from flask import Blueprint, Response, flash, redirect, render_template, request, url_for
from flask_login import current_user
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.extensions import db
from app.models import Attendance, AttendanceSetting, LeaveRequest, Role, User, Zone
from app.utils.audit import log_action
from app.utils.decorators import role_required
from app.utils.time_utils import now_in_tz, parse_time


admin_bp = Blueprint("admin", __name__, url_prefix="/admin")


def _today_in_system_tz() -> date:
    setting = AttendanceSetting.get_active()
    tz_name = setting.timezone if setting else "Asia/Makassar"
    return now_in_tz(tz_name).date()


def _resolve_period(period: str, anchor_date: date) -> tuple[date, date]:
    if period == "week":
        start_date = anchor_date - timedelta(days=anchor_date.weekday())
        end_date = start_date + timedelta(days=6)
        return start_date, end_date

    if period == "month":
        start_date = anchor_date.replace(day=1)
        next_month = (start_date.replace(day=28) + timedelta(days=4)).replace(day=1)
        end_date = next_month - timedelta(days=1)
        return start_date, end_date

    return anchor_date, anchor_date


def _resolve_report_filters():
    period = request.args.get("period", "day").lower()
    if period not in {"day", "week", "month"}:
        period = "day"

    raw_anchor = request.args.get("anchor", "").strip()
    if raw_anchor:
        try:
            anchor_date = datetime.strptime(raw_anchor, "%Y-%m-%d").date()
        except ValueError:
            anchor_date = _today_in_system_tz()
    else:
        anchor_date = _today_in_system_tz()

    start_date, end_date = _resolve_period(period, anchor_date)
    status_filter = request.args.get("status", "").strip().upper()
    institution_filter = request.args.get("institution", "").strip()
    department_filter = request.args.get("department", "").strip()
    user_id_filter = request.args.get("user_id", "").strip()

    return {
        "period": period,
        "anchor_date": anchor_date,
        "start_date": start_date,
        "end_date": end_date,
        "status": status_filter,
        "institution": institution_filter,
        "department": department_filter,
        "user_id": user_id_filter,
    }


def _build_admin_report_query(filters):
    query = Attendance.query.join(User).filter(
        User.deleted_at.is_(None),
        Attendance.date >= filters["start_date"],
        Attendance.date <= filters["end_date"],
    )

    if filters["status"]:
        query = query.filter(Attendance.status == filters["status"])
    if filters["institution"]:
        query = query.filter(User.institution.ilike(f"%{filters['institution']}%"))
    if filters["department"]:
        query = query.filter(User.department.ilike(f"%{filters['department']}%"))
    if filters["user_id"].isdigit():
        query = query.filter(User.id == int(filters["user_id"]))

    return query


@admin_bp.route("/dashboard")
@role_required("ADMIN")
def dashboard():
    total_users = User.query.filter(User.deleted_at.is_(None)).count()
    total_intern = User.query.join(Role).filter(Role.name == "INTERN", User.deleted_at.is_(None)).count()
    total_staff = User.query.join(Role).filter(Role.name == "STAFF", User.deleted_at.is_(None)).count()
    total_admin = User.query.join(Role).filter(Role.name == "ADMIN", User.deleted_at.is_(None)).count()

    today = _today_in_system_tz()
    attendance_today = Attendance.query.filter_by(date=today).count()
    late_today = Attendance.query.filter_by(date=today, status="TERLAMBAT").count()
    pending_leave = LeaveRequest.query.filter_by(status="PENDING").count()

    return render_template(
        "admin/dashboard.html",
        total_users=total_users,
        total_intern=total_intern,
        total_staff=total_staff,
        total_admin=total_admin,
        attendance_today=attendance_today,
        late_today=late_today,
        pending_leave=pending_leave,
    )


@admin_bp.route("/users")
@role_required("ADMIN")
def users():
    role_filter = request.args.get("role")
    query = User.query.filter(User.deleted_at.is_(None))

    if role_filter:
        query = query.join(Role).filter(Role.name == role_filter)

    users_data = query.order_by(User.id.desc()).all()
    roles = Role.query.order_by(Role.name.asc()).all()
    return render_template("admin/users.html", users=users_data, roles=roles)


@admin_bp.route("/users/create", methods=["GET", "POST"])
@role_required("ADMIN")
def create_user():
    roles = Role.query.order_by(Role.name.asc()).all()
    if request.method == "POST":
        role_id = int(request.form.get("role_id"))

        user = User(
            name=request.form.get("name", "").strip(),
            email=request.form.get("email", "").strip().lower(),
            username=request.form.get("username", "").strip().lower(),
            role_id=role_id,
            identity_number=request.form.get("identity_number", "").strip() or None,
            phone=request.form.get("phone", "").strip() or None,
            institution=request.form.get("institution", "").strip() or None,
            department=request.form.get("department", "").strip() or None,
            is_active=request.form.get("is_active") == "on",
        )
        user.set_password(request.form.get("password", "123456"))

        db.session.add(user)
        db.session.commit()
        log_action(
            actor_user_id=current_user.id,
            action="CREATE_USER",
            entity="users",
            entity_id=user.id,
            metadata={"role_id": user.role_id, "email": user.email, "username": user.username},
        )
        db.session.commit()

        flash("User berhasil dibuat", "success")
        return redirect(url_for("admin.users"))

    return render_template("admin/user_form.html", roles=roles, user=None)


@admin_bp.route("/users/<int:user_id>/edit", methods=["GET", "POST"])
@role_required("ADMIN")
def edit_user(user_id: int):
    user = User.query.get_or_404(user_id)
    roles = Role.query.order_by(Role.name.asc()).all()

    if request.method == "POST":
        user.name = request.form.get("name", "").strip()
        user.email = request.form.get("email", "").strip().lower()
        user.username = request.form.get("username", "").strip().lower()
        user.role_id = int(request.form.get("role_id"))
        user.identity_number = request.form.get("identity_number", "").strip() or None
        user.phone = request.form.get("phone", "").strip() or None
        user.institution = request.form.get("institution", "").strip() or None
        user.department = request.form.get("department", "").strip() or None
        user.is_active = request.form.get("is_active") == "on"

        new_password = request.form.get("password", "").strip()
        if new_password:
            user.set_password(new_password)

        log_action(
            actor_user_id=current_user.id,
            action="UPDATE_USER",
            entity="users",
            entity_id=user.id,
            metadata={"role_id": user.role_id, "is_active": user.is_active},
        )
        db.session.commit()
        flash("User berhasil diperbarui", "success")
        return redirect(url_for("admin.users"))

    return render_template("admin/user_form.html", roles=roles, user=user)


@admin_bp.route("/users/<int:user_id>/delete", methods=["POST"])
@role_required("ADMIN")
def delete_user(user_id: int):
    user = User.query.get_or_404(user_id)
    user.deleted_at = datetime.utcnow()
    user.is_active = False
    log_action(
        actor_user_id=current_user.id,
        action="SOFT_DELETE_USER",
        entity="users",
        entity_id=user.id,
        metadata={"email": user.email},
    )
    db.session.commit()
    flash("User dihapus secara soft delete", "warning")
    return redirect(url_for("admin.users"))


@admin_bp.route("/zones")
@role_required("ADMIN")
def zones():
    zone_items = Zone.query.order_by(Zone.id.desc()).all()
    return render_template("admin/zones.html", zones=zone_items)


@admin_bp.route("/zones/create", methods=["POST"])
@role_required("ADMIN")
def create_zone():
    zone = Zone(
        name=request.form.get("name", "").strip(),
        description=request.form.get("description", "").strip() or None,
        latitude=float(request.form.get("latitude", "0")),
        longitude=float(request.form.get("longitude", "0")),
        radius=float(request.form.get("radius", "100")),
        is_active=request.form.get("is_active") == "on",
    )
    db.session.add(zone)
    db.session.commit()
    log_action(
        actor_user_id=current_user.id,
        action="CREATE_ZONE",
        entity="zones",
        entity_id=zone.id,
        metadata={"name": zone.name, "radius": zone.radius, "is_active": zone.is_active},
    )
    db.session.commit()
    flash("Zona berhasil ditambahkan", "success")
    return redirect(url_for("admin.zones"))


@admin_bp.route("/zones/<int:zone_id>/edit", methods=["GET", "POST"])
@role_required("ADMIN")
def edit_zone(zone_id: int):
    zone = Zone.query.get_or_404(zone_id)

    if request.method == "POST":
        zone.name = request.form.get("name", "").strip()
        zone.description = request.form.get("description", "").strip() or None
        zone.latitude = float(request.form.get("latitude", "0"))
        zone.longitude = float(request.form.get("longitude", "0"))
        zone.radius = float(request.form.get("radius", "100"))
        zone.is_active = request.form.get("is_active") == "on"

        log_action(
            actor_user_id=current_user.id,
            action="UPDATE_ZONE",
            entity="zones",
            entity_id=zone.id,
            metadata={"name": zone.name, "radius": zone.radius, "is_active": zone.is_active},
        )
        db.session.commit()
        flash("Zona berhasil diperbarui", "success")
        return redirect(url_for("admin.zones"))

    return render_template("admin/zone_form.html", zone=zone)


@admin_bp.route("/zones/<int:zone_id>/delete", methods=["POST"])
@role_required("ADMIN")
def delete_zone(zone_id: int):
    zone = Zone.query.get_or_404(zone_id)

    used_count = Attendance.query.filter_by(zone_id=zone.id).count()
    if used_count > 0:
        flash("Zona tidak bisa dihapus karena sudah dipakai data absensi", "danger")
        return redirect(url_for("admin.zones"))

    log_action(
        actor_user_id=current_user.id,
        action="DELETE_ZONE",
        entity="zones",
        entity_id=zone.id,
        metadata={"name": zone.name},
    )
    db.session.delete(zone)
    db.session.commit()
    flash("Zona berhasil dihapus", "success")
    return redirect(url_for("admin.zones"))


@admin_bp.route("/zones/<int:zone_id>/toggle", methods=["POST"])
@role_required("ADMIN")
def toggle_zone(zone_id: int):
    zone = Zone.query.get_or_404(zone_id)
    zone.is_active = not zone.is_active
    log_action(
        actor_user_id=current_user.id,
        action="TOGGLE_ZONE",
        entity="zones",
        entity_id=zone.id,
        metadata={"is_active": zone.is_active},
    )
    db.session.commit()
    flash("Status zona diperbarui", "success")
    return redirect(url_for("admin.zones"))


@admin_bp.route("/settings", methods=["GET", "POST"])
@role_required("ADMIN")
def settings():
    setting = AttendanceSetting.get_active()

    if request.method == "POST":
        check_in_time = parse_time(request.form.get("check_in_time"), setting.check_in_time)
        check_out_time = parse_time(request.form.get("check_out_time"), setting.check_out_time)
        tolerance = int(request.form.get("late_tolerance_minutes", setting.late_tolerance_minutes))
        timezone = request.form.get("timezone", "Asia/Makassar").strip() or "Asia/Makassar"

        setting.check_in_time = check_in_time
        setting.check_out_time = check_out_time
        setting.late_tolerance_minutes = tolerance
        setting.timezone = timezone

        log_action(
            actor_user_id=current_user.id,
            action="UPDATE_ATTENDANCE_SETTINGS",
            entity="attendance_settings",
            entity_id=setting.id,
            metadata={
                "check_in_time": setting.check_in_time.strftime("%H:%M"),
                "check_out_time": setting.check_out_time.strftime("%H:%M"),
                "late_tolerance_minutes": setting.late_tolerance_minutes,
                "timezone": setting.timezone,
            },
        )
        db.session.commit()
        flash("Pengaturan absensi diperbarui", "success")
        return redirect(url_for("admin.settings"))

    return render_template("admin/settings.html", setting=setting)


@admin_bp.route("/attendance")
@role_required("ADMIN")
def attendance():
    selected_date = request.args.get("date")
    query = Attendance.query.join(User).filter(User.deleted_at.is_(None))
    if selected_date:
        query = query.filter(Attendance.date == datetime.strptime(selected_date, "%Y-%m-%d").date())
    rows = query.order_by(Attendance.date.desc(), Attendance.id.desc()).all()
    return render_template("admin/attendance.html", attendances=rows)


@admin_bp.route("/leave-requests")
@role_required("ADMIN")
def leave_requests():
    rows = LeaveRequest.query.order_by(LeaveRequest.id.desc()).all()
    return render_template("admin/leave_requests.html", leave_requests=rows)


@admin_bp.route("/reports")
@role_required("ADMIN")
def reports():
    filters = _resolve_report_filters()
    rows = _build_admin_report_query(filters).order_by(Attendance.date.desc(), Attendance.id.desc()).all()
    users = User.query.filter(User.deleted_at.is_(None)).order_by(User.name.asc()).all()
    return render_template("admin/reports.html", attendances=rows, filters=filters, users=users)


@admin_bp.route("/reports/export-csv")
@role_required("ADMIN")
def export_reports_csv():
    filters = _resolve_report_filters()
    rows = _build_admin_report_query(filters).order_by(Attendance.date.desc(), Attendance.id.desc()).all()

    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(["Tanggal", "Nama", "Masuk", "Pulang", "Status", "Zona"])

    for row in rows:
        writer.writerow(
            [
                row.date.isoformat(),
                row.user.name,
                row.check_in.strftime("%H:%M") if row.check_in else "-",
                row.check_out.strftime("%H:%M") if row.check_out else "-",
                row.status,
                row.zone.name if row.zone else "-",
            ]
        )

    return Response(
        output.getvalue(),
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=laporan_kehadiran.csv"},
    )


@admin_bp.route("/reports/export-excel")
@role_required("ADMIN")
def export_reports_excel():
    filters = _resolve_report_filters()
    rows = _build_admin_report_query(filters).order_by(Attendance.date.asc(), Attendance.id.asc()).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Laporan Admin"

    ws.merge_cells("A1:H1")
    ws["A1"] = "Laporan Kehadiran Admin"
    ws["A1"].font = Font(size=14, bold=True)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:H2")
    ws["A2"] = (
        f"Periode: {filters['start_date'].isoformat()} s/d {filters['end_date'].isoformat()} "
        f"({filters['period'].upper()})"
    )
    ws["A2"].alignment = Alignment(horizontal="center")

    headers = ["No", "Tanggal", "Nama", "Instansi", "Divisi/Prodi", "Masuk", "Pulang", "Status"]
    for idx, label in enumerate(headers, start=1):
        cell = ws.cell(row=4, column=idx, value=label)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
        cell.alignment = Alignment(horizontal="center")

    for i, row in enumerate(rows, start=1):
        ws.cell(row=4 + i, column=1, value=i)
        ws.cell(row=4 + i, column=2, value=row.date.isoformat())
        ws.cell(row=4 + i, column=3, value=row.user.name)
        ws.cell(row=4 + i, column=4, value=row.user.institution or "-")
        ws.cell(row=4 + i, column=5, value=row.user.department or "-")
        ws.cell(row=4 + i, column=6, value=row.check_in.strftime("%H:%M") if row.check_in else "-")
        ws.cell(row=4 + i, column=7, value=row.check_out.strftime("%H:%M") if row.check_out else "-")
        ws.cell(row=4 + i, column=8, value=row.status)

    for col, width in {"A": 6, "B": 14, "C": 26, "D": 24, "E": 18, "F": 12, "G": 12, "H": 14}.items():
        ws.column_dimensions[col].width = width

    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    filename = f"laporan_admin_{filters['period']}_{filters['start_date']}_{filters['end_date']}.xlsx"
    return Response(
        out.getvalue(),
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@admin_bp.route("/reports/export-pdf")
@role_required("ADMIN")
def export_reports_pdf():
    filters = _resolve_report_filters()
    rows = _build_admin_report_query(filters).order_by(Attendance.date.asc(), Attendance.id.asc()).all()

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=20)
    styles = getSampleStyleSheet()
    story = [
        Paragraph("Laporan Kehadiran Admin", styles["Title"]),
        Paragraph(
            (
                f"Periode: {filters['start_date'].isoformat()} s/d {filters['end_date'].isoformat()} "
                f"({filters['period'].upper()})"
            ),
            styles["Normal"],
        ),
        Spacer(1, 10),
    ]

    data = [["No", "Tanggal", "Nama", "Instansi", "Divisi/Prodi", "Masuk", "Pulang", "Status"]]
    for i, row in enumerate(rows, start=1):
        data.append(
            [
                str(i),
                row.date.isoformat(),
                row.user.name,
                row.user.institution or "-",
                row.user.department or "-",
                row.check_in.strftime("%H:%M") if row.check_in else "-",
                row.check_out.strftime("%H:%M") if row.check_out else "-",
                row.status,
            ]
        )

    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0F172A")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B0B7C3")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.whitesmoke, colors.HexColor("#EEF3FB")]),
            ]
        )
    )
    story.append(table)
    doc.build(story)

    buffer.seek(0)
    filename = f"laporan_admin_{filters['period']}_{filters['start_date']}_{filters['end_date']}.pdf"
    return Response(
        buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
